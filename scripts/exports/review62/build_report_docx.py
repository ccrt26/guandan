"""Convert the reviewed Markdown report to a plain, readable Word report."""
from pathlib import Path
import re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = next(p for p in Path(__file__).resolve().parents if (p/'packages').is_dir() and (p/'package.json').is_file())
sys.path.insert(0, str(ROOT/'scripts'))
from project_paths import delivery_output
OUT = delivery_output()
OUT.mkdir(parents=True, exist_ok=True)
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.8)
section.left_margin=section.right_margin=Inches(1)
for name,size in [('Normal',11),('Title',21),('Heading 1',15),('Heading 2',12.5)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    fonts=s.element.get_or_add_rPr().rFonts
    for key in list(fonts.attrib):
        if key.endswith('Theme')or key.endswith('theme'):del fonts.attrib[key]
    fonts.set(qn('w:eastAsia'),'Heiti SC')
    for border in list(s.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.line_spacing=1.2
    if name!='Normal':s.paragraph_format.keep_with_next=True
doc.core_properties.title='掼蛋本地AI训练方案与执行复查报告'
doc.core_properties.author='REVIEW6.2资料整理'

def inline(p,text,size=None):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        if not token:continue
        bold=token.startswith('**');code=token.startswith('`')
        r=p.add_run(token[2:-2]if bold else token[1:-1]if code else token)
        r.bold=bold
        if size:r.font.size=Pt(size)
        if code:r.font.size=Pt(10)

def make_table(rows):
    count=len(rows[0]);table=doc.add_table(rows=1, cols=count);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    if rows[0][0]=='阶段':widths=[.65,2.0,3.85]
    elif count==4:widths=[2.6,1.3,1.3,1.3]
    elif count==2:widths=[1.6,4.9]
    elif rows[0][0]=='项目':widths=[1.6,1.1,3.8]
    else:widths=[1.0,2.75,2.75]
    for i,c in enumerate(table.columns):c.width=Inches(widths[i])
    for ri,row in enumerate(rows):
        cells=table.rows[0].cells if ri==0 else table.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for i,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,n in [('top',85),('bottom',85),('left',100),('right',100)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if ri==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF3');tcpr.append(shading)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.13
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value,10.5)
            if ri==0:
                for r in p.runs:r.bold=True
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(OUT/'综合复查报告.md').read_text().splitlines();i=0
while i<len(lines):
    s=lines[i].strip();i+=1
    if not s:continue
    if s.startswith('|'):
        rows=[]
        while True:
            if not re.match(r'^\|[\s:|\-]+\|$',s):rows.append([v.strip()for v in s.strip('|').split('|')])
            if i>=len(lines)or not lines[i].strip().startswith('|'):break
            s=lines[i].strip();i+=1
        make_table(rows);continue
    if s.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,s[2:])
    elif s.startswith('## '):p=doc.add_paragraph(style='Heading 1');inline(p,s[3:])
    elif s.startswith('### '):p=doc.add_paragraph(style='Heading 2');inline(p,s[4:])
    else:p=doc.add_paragraph();inline(p,s)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
path=OUT/'综合复查报告.docx';doc.save(path)
print(path)
